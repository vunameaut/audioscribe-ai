import json
from pathlib import Path
from typing import Dict, Any, List
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    """Set background color of a table cell in python-docx."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

class ExportService:
    @staticmethod
    def export_to_docx(data: Dict[str, Any], output_path: Path):
        doc = Document()

        # Styles
        normal_style = doc.styles['Normal']
        normal_style.font.name = 'Calibri'
        normal_style.font.size = Pt(11)
        normal_style.font.color.rgb = RGBColor(0x22, 0x22, 0x22)

        summary = data.get("summary", {})
        transcript = data.get("transcript", [])
        speakers = data.get("speakers", [])

        # Document Title
        title_p = doc.add_paragraph()
        title_run = title_p.add_run(summary.get("title", "BIÊN BẢN VÀ BẢN GỠ BĂNG CUỘC HỌP"))
        title_run.font.size = Pt(20)
        title_run.font.bold = True
        title_run.font.color.rgb = RGBColor(0x1e, 0x29, 0x3b)
        title_p.paragraph_format.space_after = Pt(4)

        # Subtitle / Date
        sub_p = doc.add_paragraph()
        sub_run = sub_p.add_run("Tài liệu được tổng hợp tự động bởi AudioScribe AI")
        sub_run.font.size = Pt(10)
        sub_run.font.italic = True
        sub_run.font.color.rgb = RGBColor(0x64, 0x74, 0x8b)
        sub_p.paragraph_format.space_after = Pt(16)

        # 1. TỔNG QUAN
        doc.add_heading("1. TỔNG QUAN NỘI DUNG", level=1)
        overview_p = doc.add_paragraph(summary.get("overview", "Không có thông tin tổng quan."))
        overview_p.paragraph_format.space_after = Pt(12)

        # 2. CÁC QUYẾT ĐỊNH / KẾT LUẬN
        decisions = summary.get("decisions", [])
        if decisions:
            doc.add_heading("2. CÁC QUYẾT ĐỊNH & THỐNG NHẤT CHÍNH", level=1)
            for dec in decisions:
                p = doc.add_paragraph(style='List Bullet')
                r = p.add_run(dec)
                r.bold = True
            doc.add_paragraph().paragraph_format.space_after = Pt(8)

        # 3. DANH SÁCH VIỆC CẦN LÀM (ACTION ITEMS)
        action_items = summary.get("action_items", [])
        if action_items:
            doc.add_heading("3. DANH SÁCH VIỆC CẦN LÀM (ACTION ITEMS)", level=1)
            table = doc.add_table(rows=1, cols=3)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            hdr_cells = table.rows[0].cells
            hdr_cells[0].text = "Người phụ trách"
            hdr_cells[1].text = "Nhiệm vụ cần thực hiện"
            hdr_cells[2].text = "Thời hạn / Ghi chú"

            for cell in hdr_cells:
                set_cell_background(cell, "F1F5F9")
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.bold = True
                        run.font.size = Pt(10)

            for item in action_items:
                row_cells = table.add_row().cells
                row_cells[0].text = str(item.get("assignee", "-"))
                row_cells[1].text = str(item.get("task", "-"))
                row_cells[2].text = str(item.get("deadline", "-"))

            doc.add_paragraph().paragraph_format.space_after = Pt(14)

        # 4. QUAN ĐIỂM TỪNG NGƯỜI NÓI
        spk_viewpoints = summary.get("speaker_viewpoints", [])
        if spk_viewpoints:
            doc.add_heading("4. QUAN ĐIỂM & ĐÓNG GÓP THEO TỪNG THÀNH VIÊN", level=1)
            for vp in spk_viewpoints:
                p = doc.add_paragraph()
                r_name = p.add_run(f"• {vp.get('speaker', 'Người nói')}: ")
                r_name.bold = True
                p.add_run(vp.get('role_or_perspective', ''))
                stmts = vp.get("key_statements", [])
                for st in stmts:
                    st_p = doc.add_paragraph(style='List Bullet 2')
                    st_p.add_run(st)

        # 5. BẢN GỠ BĂNG CHI TIẾT
        doc.add_heading("5. BẢN GỠ BĂNG LỜI THOẠI ĐẦY ĐỦ", level=1)
        for item in transcript:
            speaker = item.get("speaker", "Người nói")
            start = item.get("start", "00:00:00")
            end = item.get("end", "")
            time_str = f"[{start} - {end}]" if end else f"[{start}]"
            text = item.get("text", "")

            p = doc.add_paragraph()
            r_spk = p.add_run(f"{speaker} ")
            r_spk.bold = True
            r_spk.font.color.rgb = RGBColor(0x3b, 0x82, 0xf6)

            r_time = p.add_run(f"{time_str}: ")
            r_time.font.size = Pt(9)
            r_time.font.color.rgb = RGBColor(0x94, 0xa3, 0xb8)

            p.add_run(text)
            p.paragraph_format.space_after = Pt(6)

        doc.save(str(output_path))
        return output_path

    @staticmethod
    def export_to_markdown(data: Dict[str, Any], output_path: Path):
        summary = data.get("summary", {})
        transcript = data.get("transcript", [])

        md = []
        md.append(f"# {summary.get('title', 'Biên bản & Gỡ băng cuộc trò chuyện')}\n")
        md.append("> *Tự động tổng hợp bởi AudioScribe AI*\n")

        md.append("## 1. Tổng quan nội dung")
        md.append(summary.get("overview", "") + "\n")

        decisions = summary.get("decisions", [])
        if decisions:
            md.append("## 2. Các quyết định & Thống nhất chính")
            for d in decisions:
                md.append(f"- **{d}**")
            md.append("")

        action_items = summary.get("action_items", [])
        if action_items:
            md.append("## 3. Danh sách việc cần làm (Action Items)")
            md.append("| Người phụ trách | Nhiệm vụ | Thời hạn |")
            md.append("| :--- | :--- | :--- |")
            for item in action_items:
                md.append(f"| {item.get('assignee', '-')} | {item.get('task', '-')} | {item.get('deadline', '-')} |")
            md.append("")

        spk_viewpoints = summary.get("speaker_viewpoints", [])
        if spk_viewpoints:
            md.append("## 4. Quan điểm & Đóng góp từng người nói")
            for vp in spk_viewpoints:
                md.append(f"### {vp.get('speaker')}")
                md.append(f"*{vp.get('role_or_perspective', '')}*")
                for st in vp.get("key_statements", []):
                    md.append(f"- {st}")
                md.append("")

        md.append("## 5. Bản gỡ băng chi tiết")
        for item in transcript:
            md.append(f"**{item.get('speaker')}** `[{item.get('start')} -> {item.get('end')}]`:")
            md.append(f"> {item.get('text')}\n")

        output_path.write_text("\n".join(md), encoding="utf-8")
        return output_path

    @staticmethod
    def export_to_txt(data: Dict[str, Any], output_path: Path):
        summary = data.get("summary", {})
        transcript = data.get("transcript", [])

        lines = []
        lines.append(f"=== {summary.get('title', 'BẢN GỠ BĂNG CUỘC HỌP')} ===")
        lines.append(f"\nTỔNG QUAN:\n{summary.get('overview', '')}\n")

        decisions = summary.get("decisions", [])
        if decisions:
            lines.append("CÁC QUYẾT ĐỊNH CHÍNH:")
            for d in decisions:
                lines.append(f"- {d}")
            lines.append("")

        lines.append("CHI TIẾT LỜI THOẠI:")
        lines.append("-" * 50)
        for item in transcript:
            lines.append(f"[{item.get('start')} - {item.get('end')}] {item.get('speaker')}: {item.get('text')}")

        output_path.write_text("\n".join(lines), encoding="utf-8")
        return output_path

    @staticmethod
    def export_to_srt(data: Dict[str, Any], output_path: Path):
        transcript = data.get("transcript", [])

        def to_srt_time(t_str: str) -> str:
            parts = t_str.strip().split(":")
            if len(parts) == 2:
                m, s = int(parts[0]), float(parts[1])
                h = 0
            elif len(parts) == 3:
                h, m, s = int(parts[0]), int(parts[1]), float(parts[2])
            else:
                return "00:00:00,000"
            sec = int(s)
            ms = int(round((s - sec) * 1000))
            return f"{h:02d}:{m:02d}:{sec:02d},{ms:03d}"

        srt_lines = []
        for idx, item in enumerate(transcript, start=1):
            start = to_srt_time(item.get("start", "00:00"))
            end = to_srt_time(item.get("end", item.get("start", "00:00")))
            text = f"[{item.get('speaker', '')}]: {item.get('text', '')}"

            srt_lines.append(str(idx))
            srt_lines.append(f"{start} --> {end}")
            srt_lines.append(text)
            srt_lines.append("")

        output_path.write_text("\n".join(srt_lines), encoding="utf-8")
        return output_path
